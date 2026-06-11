from __future__ import annotations

import io
import json
import threading
import uuid
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional

from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, JSONResponse
from fastapi.staticfiles import StaticFiles
from pptx import Presentation
from pptx.util import Inches
from pydantic import BaseModel

from config import (
    BUILTIN_TEMPLATE_PATHS,
    INPUTS_DIR,
    OUTPUTS_DIR,
    SAMPLE_TEXT,
    SETTINGS_PATH,
    TEMPLATES_DIR,
    TOTAL_STEPS,
    ensure_runtime_dirs,
    load_runtime_settings,
    resolve_llm_provider,
    save_runtime_settings,
)
from deps import check_dependencies
from eval_runner import run_batch_evaluation
from pipeline import TaskPipeline
from task_store import TaskStore

ensure_runtime_dirs()

app = FastAPI(title="Paper2PPT API", version="2.0.0")
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

STORE = TaskStore()
TASK_LOCK = threading.Lock()
TASKS: Dict[str, Dict[str, Any]] = STORE.load_tasks_into_memory()


def _now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def _persist_task(task_id: str, updates: Optional[Dict[str, Any]] = None) -> None:
    with TASK_LOCK:
        if updates:
            task = TASKS.get(task_id, {})
            task.update(updates)
            task["updated_at"] = _now_iso()
            TASKS[task_id] = task
        if task_id in TASKS:
            STORE.save_task(task_id, TASKS[task_id])


def _on_pipeline_update(task_id: str, updates: Dict[str, Any]) -> None:
    _persist_task(task_id, updates)


PIPELINE = TaskPipeline(store=STORE, on_update=_on_pipeline_update)


def _create_builtin_template(path: Path, width_in: float, height_in: float) -> None:
    prs = Presentation()
    prs.slide_width = Inches(width_in)
    prs.slide_height = Inches(height_in)
    title_slide = prs.slides.add_slide(prs.slide_layouts[0])
    title_slide.shapes.title.text = "Paper2PPT Built-in Template"
    subtitle = title_slide.placeholders[1]
    subtitle.text = f"{path.stem}"
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(path))


def _ensure_builtin_templates() -> None:
    if not BUILTIN_TEMPLATE_PATHS["学术极简模板"].exists():
        _create_builtin_template(BUILTIN_TEMPLATE_PATHS["学术极简模板"], 13.333, 7.5)
    if not BUILTIN_TEMPLATE_PATHS["会议汇报模板"].exists():
        _create_builtin_template(BUILTIN_TEMPLATE_PATHS["会议汇报模板"], 13.333, 8.0)
    if not BUILTIN_TEMPLATE_PATHS["课题组周报模板"].exists():
        _create_builtin_template(BUILTIN_TEMPLATE_PATHS["课题组周报模板"], 10.0, 7.5)


_ensure_builtin_templates()


class ConfirmPayload(BaseModel):
    action: str
    feedback: str = ""
    layout_overrides: Dict[str, str] = {}
    outline_overrides: Dict[str, str] = {}
    title_override: str = ""


class LLMSettingsPayload(BaseModel):
    provider: str = "dashscope"
    api_key: str = ""
    base_url: str = "https://dashscope.aliyuncs.com/compatible-mode/v1"
    model: str = "qwen3.7-plus"
    enable_thinking: bool = False


def _mask_api_key(api_key: str) -> str:
    if not api_key:
        return ""
    if len(api_key) <= 8:
        return "****"
    return f"{api_key[:4]}****{api_key[-4:]}"


def _run_worker(task_id: str) -> None:
    try:
        PIPELINE.run_until_pause_or_done(task_id)
    except Exception as exc:
        STORE.append_log(task_id, f"后台任务异常: {exc}", level="error")
        _persist_task(
            task_id,
            {
                "status": "failed",
                "message": f"后台任务异常: {exc}",
                "error": str(exc),
            },
        )


@app.on_event("startup")
def _resume_unfinished_tasks() -> None:
    for task_id, task in list(TASKS.items()):
        if task.get("status") not in {"queued", "running"}:
            continue
        if task.get("current_stage") in {"outline_confirm", "slides_confirm", "layout_confirm"}:
            continue
        STORE.append_log(task_id, "Resuming unfinished task")
        threading.Thread(target=_run_worker, args=(task_id,), daemon=True).start()


@app.get("/api/health")
def health() -> Dict[str, str]:
    deps = check_dependencies()
    return {
        "status": "ok",
        "version": "2.0.0",
        "ocr_ready": str(deps["ocr_ready"]).lower(),
        "figures_ready": str(deps["figures_ready"]).lower(),
        "tables_ready": str(deps["tables_ready"]).lower(),
    }


@app.get("/api/deps")
def deps_status() -> JSONResponse:
    return JSONResponse(check_dependencies())


@app.get("/api/settings/llm")
def get_llm_settings() -> JSONResponse:
    llm_settings = (load_runtime_settings().get("llm") or {})
    api_key = str(llm_settings.get("api_key") or "")
    return JSONResponse(
        {
            "provider": llm_settings.get("provider") or "dashscope",
            "base_url": llm_settings.get("base_url") or "https://dashscope.aliyuncs.com/compatible-mode/v1",
            "model": llm_settings.get("model") or "qwen3.7-plus",
            "enable_thinking": bool(llm_settings.get("enable_thinking", False)),
            "api_key_set": bool(api_key),
            "api_key_mask": _mask_api_key(api_key),
            "settings_path": str(SETTINGS_PATH),
        }
    )


@app.put("/api/settings/llm")
def update_llm_settings(payload: LLMSettingsPayload) -> JSONResponse:
    settings = load_runtime_settings()
    current_llm = dict(settings.get("llm") or {})
    provider = payload.provider.strip().lower() or "dashscope"
    if provider not in {"dashscope", "openai"}:
        raise HTTPException(status_code=400, detail="provider 必须是 dashscope 或 openai")
    api_key = payload.api_key.strip() or str(current_llm.get("api_key") or "").strip()
    base_url = payload.base_url.strip()
    model = payload.model.strip()
    if not base_url:
        base_url = "https://dashscope.aliyuncs.com/compatible-mode/v1" if provider == "dashscope" else "https://api.openai.com/v1"
    if not model:
        model = "qwen3.7-plus" if provider == "dashscope" else "gpt-4o-mini"
    settings["llm"] = {
        "provider": provider,
        "api_key": api_key,
        "base_url": base_url,
        "model": model,
        "enable_thinking": payload.enable_thinking,
    }
    save_runtime_settings(settings)
    return JSONResponse(
        {
            "status": "ok",
            "provider": provider,
            "base_url": base_url,
            "model": model,
            "enable_thinking": payload.enable_thinking,
            "api_key_set": bool(api_key),
            "api_key_mask": _mask_api_key(api_key),
        }
    )


@app.post("/api/tasks")
async def create_task(
    input_mode: str = Form("upload"),
    pasted_text: str = Form(""),
    page_count: int = Form(8),
    template_id: str = Form("tech-sharing"),
    style_id: str = Form("blue"),
    instruction: str = Form(""),
    file: Optional[UploadFile] = File(default=None),
    supplemental_figures: Optional[List[UploadFile]] = File(default=None),
    supplemental_figure_meta: str = Form(""),
    template_file: Optional[UploadFile] = File(default=None),
) -> Dict[str, str]:
    task_id = uuid.uuid4().hex
    source_path: Optional[str] = None
    source_kind = input_mode

    if input_mode == "upload":
        if file is None:
            raise HTTPException(status_code=400, detail="请上传 PDF、Markdown 或文本文件")
        suffix = Path(file.filename or "paper.pdf").suffix or ".pdf"
        normalized_suffix = suffix.lower()
        if normalized_suffix not in {".pdf", ".md", ".markdown", ".txt"}:
            raise HTTPException(status_code=400, detail="仅支持上传 PDF、Markdown（.md/.markdown）或文本（.txt）文件")
        input_path = INPUTS_DIR / f"{task_id}{suffix}"
        input_path.write_bytes(await file.read())
        source_path = str(input_path)
        source_kind = "pdf" if normalized_suffix == ".pdf" else "text"
    elif input_mode == "paste":
        text = pasted_text.strip()
        if not text:
            raise HTTPException(status_code=400, detail="粘贴文本不能为空")
        input_path = INPUTS_DIR / f"{task_id}.txt"
        input_path.write_text(text, encoding="utf-8")
        source_path = str(input_path)
        source_kind = "text"
    elif input_mode == "example":
        source_path = str(SAMPLE_TEXT)
        source_kind = "text"
    else:
        raise HTTPException(status_code=400, detail="未知输入模式")

    supplemental_figure_paths: List[Dict[str, str]] = []
    figure_meta: List[Dict[str, str]] = []
    if supplemental_figure_meta.strip():
        try:
            loaded_meta = json.loads(supplemental_figure_meta)
            if isinstance(loaded_meta, list):
                figure_meta = [item for item in loaded_meta if isinstance(item, dict)]
        except json.JSONDecodeError:
            figure_meta = []
    if supplemental_figures:
        figure_dir = OUTPUTS_DIR / task_id / "user_figures"
        figure_dir.mkdir(parents=True, exist_ok=True)
        for idx, figure_file in enumerate(supplemental_figures[:12], start=1):
            meta = figure_meta[idx - 1] if idx - 1 < len(figure_meta) else {}
            suffix = Path(figure_file.filename or f"figure-{idx}.png").suffix.lower()
            if suffix not in {".png", ".jpg", ".jpeg", ".webp"}:
                raise HTTPException(status_code=400, detail="补充图片仅支持 PNG/JPG/WebP")
            figure_path = figure_dir / f"user-figure-{idx}{suffix}"
            figure_path.write_bytes(await figure_file.read())
            supplemental_figure_paths.append(
                {
                    "path": str(figure_path),
                    "kind": str(meta.get("kind") or "other"),
                    "description": str(meta.get("description") or "").strip(),
                    "filename": str(meta.get("filename") or figure_file.filename or f"figure-{idx}"),
                }
            )

    template_path: Optional[str] = None
    if template_file is not None:
        suffix = Path(template_file.filename or "template.pptx").suffix or ".pptx"
        file_path = TEMPLATES_DIR / f"{task_id}{suffix}"
        file_path.write_bytes(await template_file.read())
        template_path = str(file_path)

    llm_provider = resolve_llm_provider()
    bounded_pages = max(4, min(page_count, 20))
    task_payload = {
        "task_id": task_id,
        "status": "queued",
        "message": "任务已创建",
        "error": None,
        "created_at": _now_iso(),
        "updated_at": _now_iso(),
        "completed_steps": 0,
        "total_steps": TOTAL_STEPS,
        "current_stage": "parse",
        "timeline": [],
        "pending_confirmation": None,
        "result": None,
        "output_path": None,
        "html_path": None,
        "config": {
            "input_mode": input_mode,
            "page_count": bounded_pages,
            "template_id": template_id,
            "style_id": style_id,
            "template_path": template_path,
            "instruction": instruction.strip(),
            "llm_provider": llm_provider,
            "source_path": source_path,
            "source_kind": source_kind,
            "supplemental_figure_paths": supplemental_figure_paths,
        },
    }
    _persist_task(task_id, task_payload)
    STORE.append_log(task_id, "Task created")

    threading.Thread(target=_run_worker, args=(task_id,), daemon=True).start()
    return {"task_id": task_id}


@app.get("/api/tasks/{task_id}")
def get_task(task_id: str) -> Dict[str, Any]:
    with TASK_LOCK:
        task = TASKS.get(task_id)
    if not task:
        task = STORE.get_task(task_id)
        if task:
            with TASK_LOCK:
                TASKS[task_id] = task
    if not task:
        raise HTTPException(status_code=404, detail="任务不存在")
    return task


@app.post("/api/tasks/{task_id}/confirm")
def confirm_task(task_id: str, payload: ConfirmPayload) -> Dict[str, str]:
    with TASK_LOCK:
        task = TASKS.get(task_id)
    if not task:
        raise HTTPException(status_code=404, detail="任务不存在")
    if task.get("status") != "waiting_confirmation":
        raise HTTPException(status_code=409, detail="当前任务不在等待确认状态")

    action = payload.action.strip().lower()
    if action not in {"approve", "revise"}:
        raise HTTPException(status_code=400, detail="action 必须是 approve 或 revise")

    _persist_task(task_id, {"status": "running"})
    PIPELINE.continue_after_confirm(
        task_id,
        action,
        payload.feedback.strip(),
        payload.layout_overrides or {},
        payload.outline_overrides or {},
        payload.title_override.strip(),
    )
    return {"status": "ok"}


@app.post("/api/tasks/{task_id}/figures/{item_id}/upload")
async def upload_task_figure(task_id: str, item_id: str, file: UploadFile = File(...)) -> JSONResponse:
    with TASK_LOCK:
        task = TASKS.get(task_id)
    if not task:
        task = STORE.get_task(task_id)
    if not task:
        raise HTTPException(status_code=404, detail="任务不存在")
    if not item_id.startswith("layout-"):
        raise HTTPException(status_code=400, detail="图片只能绑定到布局确认页面")
    try:
        slide_idx = int(item_id.split("-", 1)[1])
    except ValueError as exc:
        raise HTTPException(status_code=400, detail="无效页面编号") from exc

    suffix = Path(file.filename or "figure.png").suffix.lower()
    if suffix not in {".png", ".jpg", ".jpeg", ".webp"}:
        raise HTTPException(status_code=400, detail="仅支持 PNG/JPG/WebP 图片")

    parsed = dict(task.get("parsed") or {})
    figures = list(parsed.get("figures") or [])
    figure_dir = Path(OUTPUTS_DIR) / task_id / "custom_figures"
    figure_dir.mkdir(parents=True, exist_ok=True)
    figure_path = figure_dir / f"{item_id}-{uuid.uuid4().hex[:8]}{suffix}"
    figure_path.write_bytes(await file.read())
    figures.append(
        {
            "page_index": slide_idx,
            "figure_id": f"User uploaded figure for slide {slide_idx + 1}",
            "caption": "",
            "image_path": str(figure_path),
        }
    )
    parsed["figures"] = figures

    figure_bindings = {str(k): v for k, v in dict(task.get("figure_bindings") or {}).items()}
    figure_bindings[str(slide_idx)] = len(figures) - 1
    layout_names = list(task.get("layout_names") or [])
    while slide_idx >= len(layout_names):
        layout_names.append("bullets")
    if not (layout_names[slide_idx].startswith("figure-") or layout_names[slide_idx] in {"image-left", "image-right", "image-hero"}):
        layout_names[slide_idx] = "figure-hero"
    pending_confirmation = task.get("pending_confirmation")
    if pending_confirmation and pending_confirmation.get("stage_id") == "layout_confirm":
        pending_confirmation = dict(pending_confirmation)
        items = []
        for item in pending_confirmation.get("items") or []:
            item = dict(item)
            if item.get("id") == item_id:
                item["figure_url"] = f"/api/tasks/{task_id}/figures/{len(figures) - 1}/preview"
                item["figure_caption"] = "用户上传的替换图片"
            items.append(item)
        pending_confirmation["items"] = items

    _persist_task(
        task_id,
        {
            "parsed": parsed,
            "figure_bindings": figure_bindings,
            "layout_names": layout_names,
            "pending_confirmation": pending_confirmation or task.get("pending_confirmation"),
        },
    )
    STORE.append_log(task_id, f"User uploaded replacement figure for {item_id}")
    return JSONResponse({"status": "ok", "item_id": item_id, "figure_index": len(figures) - 1})


@app.get("/api/tasks/{task_id}/figures/{figure_idx}/preview")
def preview_task_figure(task_id: str, figure_idx: int) -> FileResponse:
    with TASK_LOCK:
        task = TASKS.get(task_id)
    if not task:
        task = STORE.get_task(task_id)
    if not task:
        raise HTTPException(status_code=404, detail="任务不存在")
    figures = ((task.get("parsed") or {}).get("figures") or [])
    if figure_idx < 0 or figure_idx >= len(figures):
        raise HTTPException(status_code=404, detail="图片不存在")
    image_path = Path(figures[figure_idx].get("image_path") or "")
    if not image_path.exists() or not image_path.is_file():
        raise HTTPException(status_code=404, detail="图片文件不存在")
    return FileResponse(str(image_path))


@app.get("/api/tasks/{task_id}/download")
def download_task(task_id: str) -> FileResponse:
    task = get_task(task_id)
    if task.get("status") != "succeeded":
        raise HTTPException(status_code=409, detail="任务尚未完成")
    output_path = task.get("output_path")
    if not output_path or not Path(output_path).exists():
        raise HTTPException(status_code=404, detail="输出文件不存在")
    return FileResponse(
        output_path,
        media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        filename=f"paper2ppt-{task_id}.pptx",
    )


@app.get("/api/tasks/{task_id}/html/download")
def download_html(task_id: str) -> FileResponse:
    task = get_task(task_id)
    html_dir = Path(OUTPUTS_DIR) / task_id / "html"
    if not html_dir.exists():
        raise HTTPException(status_code=404, detail="HTML 输出不存在")
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as zf:
        for path in html_dir.rglob("*"):
            if path.is_file():
                zf.write(path, arcname=str(path.relative_to(html_dir)))
    buffer.seek(0)
    out = OUTPUTS_DIR / f"{task_id}-html.zip"
    out.write_bytes(buffer.getvalue())
    return FileResponse(out, media_type="application/zip", filename=f"paper2ppt-{task_id}-html.zip")


@app.get("/api/tasks/{task_id}/versions")
def list_versions(task_id: str) -> JSONResponse:
    if task_id not in TASKS and not STORE.get_task(task_id):
        raise HTTPException(status_code=404, detail="任务不存在")
    return JSONResponse({"versions": STORE.list_versions(task_id)})


@app.get("/api/tasks/{task_id}/versions/{version_no}")
def get_version(task_id: str, version_no: int) -> JSONResponse:
    if task_id not in TASKS and not STORE.get_task(task_id):
        raise HTTPException(status_code=404, detail="任务不存在")
    version = STORE.get_version(task_id, version_no)
    if not version:
        raise HTTPException(status_code=404, detail="版本不存在")
    return JSONResponse(version)


@app.post("/api/tasks/{task_id}/versions/{version_no}/restore")
def restore_version(task_id: str, version_no: int) -> JSONResponse:
    with TASK_LOCK:
        task = TASKS.get(task_id)
    if not task:
        task = STORE.get_task(task_id)
    if not task:
        raise HTTPException(status_code=404, detail="任务不存在")
    version = STORE.get_version(task_id, version_no)
    if not version:
        raise HTTPException(status_code=404, detail="版本不存在")
    snapshot = version.get("snapshot") or {}
    result = snapshot.get("result")
    if not result:
        raise HTTPException(status_code=409, detail="该版本不含可恢复的结果")
    updates = {
        "status": "succeeded",
        "message": f"已恢复到版本 v{version_no}",
        "result": result,
        "completed_steps": TOTAL_STEPS,
        "total_steps": TOTAL_STEPS,
        "current_stage": "done",
        "pending_confirmation": None,
    }
    _persist_task(task_id, updates)
    STORE.append_log(task_id, f"Restored version v{version_no}")
    return JSONResponse({"status": "ok", "result": result})


@app.get("/api/tasks/{task_id}/logs")
def list_logs(task_id: str) -> JSONResponse:
    if task_id not in TASKS and not STORE.get_task(task_id):
        raise HTTPException(status_code=404, detail="任务不存在")
    return JSONResponse({"logs": STORE.list_logs(task_id)})


@app.post("/api/eval/batch")
def batch_eval() -> JSONResponse:
    report = run_batch_evaluation()
    return JSONResponse(report)


html_root = OUTPUTS_DIR
if html_root.exists():
    @app.get("/api/tasks/{task_id}/html/{file_path:path}")
    def serve_html(task_id: str, file_path: str) -> FileResponse:
        base = Path(OUTPUTS_DIR) / task_id / "html"
        target = (base / file_path).resolve()
        if not str(target).startswith(str(base.resolve())):
            raise HTTPException(status_code=403, detail="Forbidden")
        if not target.exists() or not target.is_file():
            raise HTTPException(status_code=404, detail="File not found")
        return FileResponse(target)
