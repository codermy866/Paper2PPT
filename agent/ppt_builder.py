from __future__ import annotations

from pathlib import Path
from typing import Dict, List, Tuple

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from models import ParsedFigure, ParsedTable, SlideContent


TITLE_COLOR = RGBColor(33, 37, 41)
BODY_COLOR = RGBColor(55, 65, 81)
PLACEHOLDER_TITLE = 1
PLACEHOLDER_BODY = 2
PLACEHOLDER_SUBTITLE = 4
PLACEHOLDER_OBJECT = 7

STYLE_ACCENTS: dict[str, Tuple[int, int, int]] = {
    "blue": (0, 94, 184),
    "red": (196, 45, 43),
    "green": (46, 125, 50),
    "purple": (93, 63, 211),
}


def _apply_title_style(text_frame) -> None:
    paragraph = text_frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.LEFT
    if paragraph.runs:
        run = paragraph.runs[0]
        run.font.name = "Arial"
        run.font.size = Pt(28)
        run.font.bold = True
        run.font.color.rgb = TITLE_COLOR


def _apply_body_style(text_frame) -> None:
    for paragraph in text_frame.paragraphs:
        paragraph.space_after = Pt(6)
        for run in paragraph.runs:
            run.font.name = "Arial"
            run.font.size = Pt(21)
            run.font.color.rgb = BODY_COLOR


def _set_notes(slide, notes: str) -> None:
    if not notes:
        return
    try:
        notes_slide = slide.notes_slide
        text_frame = notes_slide.notes_text_frame
        text_frame.text = notes
    except Exception:
        # python-pptx notes support is limited across versions; ignore if unavailable.
        pass


def _find_placeholder_text_frame(slide, placeholder_type) -> object | None:
    # Placeholder objects have a `.text_frame` we can fill.
    for ph in slide.placeholders:
        try:
            if ph.placeholder_format.type == placeholder_type:
                return ph.text_frame
        except Exception:
            continue
    return None


def _layout_placeholder_types(layout) -> set[int]:
    types: set[int] = set()
    for ph in layout.placeholders:
        try:
            types.add(int(ph.placeholder_format.type))
        except Exception:
            continue
    return types


def _score_layout_name_for_slide_key(name: str, slide_key: str) -> int:
    if slide_key == "title":
        if "title slide" in name:
            return 10
        if "title" in name:
            return 7
        return 0
    if slide_key in ("background", "conclusion"):
        if "section" in name:
            return 6
        if "title and content" in name:
            return 4
    if slide_key in ("method", "results"):
        if "comparison" in name or "two content" in name:
            return 5
        if "content" in name:
            return 4
    if slide_key == "experiment_setup":
        if "content with caption" in name or "picture with caption" in name:
            return 5
        if "content" in name:
            return 3
    if "blank" in name:
        return -6
    if "title only" in name:
        return -2
    return 0


def _select_layout(prs: Presentation, slide_key: str):
    best_layout = None
    best_score = -10**9

    for layout in prs.slide_layouts:
        name = (layout.name or "").lower()
        types = _layout_placeholder_types(layout)
        has_title = PLACEHOLDER_TITLE in types
        has_body = PLACEHOLDER_BODY in types or PLACEHOLDER_OBJECT in types

        score = 0
        if has_title:
            score += 3
        if has_body:
            score += 3

        score += _score_layout_name_for_slide_key(name, slide_key)

        if score > best_score:
            best_score = score
            best_layout = layout

    return best_layout if best_layout is not None else prs.slide_layouts[0]


def _remove_all_slides(prs: Presentation) -> None:
    # Rebuild deck from template layouts/theme by dropping existing slide instances.
    sld_id_lst = prs.slides._sldIdLst
    for sld_id in list(sld_id_lst):
        rel_id = sld_id.rId
        prs.part.drop_rel(rel_id)
        sld_id_lst.remove(sld_id)


def _set_title_on_slide(slide, title_text: str) -> None:
    # Prefer placeholder TITLE if present; fall back to `slide.shapes.title`.
    try:
        title_tf = _find_placeholder_text_frame(slide, PLACEHOLDER_TITLE)
        if title_tf is not None:
            title_tf.text = title_text
            return
    except Exception:
        pass

    title_shape = slide.shapes.title
    title_shape.text = title_text


def _find_body_text_frame(slide):
    body_tf = _find_placeholder_text_frame(slide, PLACEHOLDER_BODY)
    if body_tf is not None:
        return body_tf

    object_tf = _find_placeholder_text_frame(slide, PLACEHOLDER_OBJECT)
    if object_tf is not None:
        return object_tf

    for shape in slide.shapes:
        if not getattr(shape, "has_text_frame", False):
            continue
        try:
            ph_type = int(shape.placeholder_format.type)
            if ph_type in (PLACEHOLDER_TITLE, PLACEHOLDER_SUBTITLE):
                continue
        except Exception:
            pass
        return shape.text_frame
    return None


def _append_figure_slides(prs: Presentation, figures: List[ParsedFigure], accent_color: RGBColor) -> None:
    for figure in figures[:4]:
        layout = prs.slide_layouts[5] if len(prs.slide_layouts) > 5 else prs.slide_layouts[0]
        slide = prs.slides.add_slide(layout)
        title = slide.shapes.title
        if title is not None:
            title.text = figure.caption or f"Figure {figure.figure_id}"
        image_path = Path(figure.image_path)
        if image_path.exists():
            slide.shapes.add_picture(
                str(image_path),
                Inches(1.0),
                Inches(1.6),
                width=Inches(11.0),
            )
        line = slide.shapes.add_shape(1, Inches(0.8), Inches(1.35), Inches(2.0), Inches(0.08))
        line.fill.solid()
        line.fill.fore_color.rgb = accent_color


def _append_table_slides(prs: Presentation, tables: List[ParsedTable], accent_color: RGBColor) -> None:
    for table in tables[:3]:
        layout = prs.slide_layouts[1] if len(prs.slide_layouts) > 1 else prs.slide_layouts[0]
        slide = prs.slides.add_slide(layout)
        title = slide.shapes.title
        if title is not None:
            title.text = f"Table {table.table_id}"
        image_path = Path(table.image_path)
        if table.image_path and image_path.exists():
            slide.shapes.add_picture(
                str(image_path),
                Inches(0.9),
                Inches(1.6),
                width=Inches(11.5),
            )
        else:
            body_tf = _find_body_text_frame(slide)
            if body_tf is not None:
                body_tf.clear()
                lines = [line.strip() for line in table.markdown.splitlines() if line.strip()][:12]
                for idx, line in enumerate(lines):
                    paragraph = body_tf.paragraphs[0] if idx == 0 else body_tf.add_paragraph()
                    paragraph.text = line
                    paragraph.level = 0
                    paragraph.space_after = Pt(6)
                _apply_body_style(body_tf)
        line = slide.shapes.add_shape(1, Inches(0.8), Inches(1.35), Inches(2.0), Inches(0.08))
        line.fill.solid()
        line.fill.fore_color.rgb = accent_color


def _add_picture_contain(slide, image_path: Path, left, top, width, height) -> None:
    if not image_path.exists():
        return
    try:
        with Image.open(image_path) as image:
            img_w, img_h = image.size
    except Exception:
        slide.shapes.add_picture(str(image_path), left, top, width=width)
        return
    if img_w <= 0 or img_h <= 0:
        return
    box_w = int(width)
    box_h = int(height)
    box_ratio = box_w / box_h
    image_ratio = img_w / img_h
    if image_ratio >= box_ratio:
        final_width = box_w
        final_height = int(box_w / image_ratio)
    else:
        final_height = box_h
        final_width = int(box_h * image_ratio)
    final_width = min(final_width, box_w)
    final_height = min(final_height, box_h)
    final_left = int(left) + max(0, int((box_w - final_width) / 2))
    final_top = int(top) + max(0, int((box_h - final_height) / 2))
    slide.shapes.add_picture(str(image_path), final_left, final_top, width=final_width, height=final_height)


def build_ppt(
    slides: List[SlideContent],
    output_path: str,
    template_path: str | None = None,
    style_key: str = "blue",
    figures: List[ParsedFigure] | None = None,
    tables: List[ParsedTable] | None = None,
    layout_names: List[str] | None = None,
    figure_bindings: Dict[int, int] | None = None,
) -> str:
    accent_rgb = STYLE_ACCENTS.get(style_key, STYLE_ACCENTS["blue"])
    accent_color = RGBColor(*accent_rgb)
    figure_items = figures or []
    table_items = tables or []
    layout_items = layout_names or []
    bound_figures = figure_bindings or {}
    if template_path:
        prs = Presentation(template_path)
        _remove_all_slides(prs)

        for index, slide_content in enumerate(slides):
            layout = _select_layout(prs, slide_content.slide_key)
            slide = prs.slides.add_slide(layout)
            _set_title_on_slide(slide, slide_content.title)

            body_tf = _find_body_text_frame(slide)

            if body_tf is not None:
                body_tf.clear()
                for bullet_index, bullet in enumerate(slide_content.bullets):
                    paragraph = body_tf.paragraphs[0] if bullet_index == 0 else body_tf.add_paragraph()
                    paragraph.text = bullet
                    paragraph.level = 0
                    paragraph.space_after = Pt(6)
                _apply_body_style(body_tf)

            figure_idx = bound_figures.get(index)
            if figure_idx is not None and 0 <= figure_idx < len(figure_items):
                image_path = Path(figure_items[figure_idx].image_path)
                if image_path.exists():
                    _add_picture_contain(slide, image_path, Inches(7.15), Inches(1.45), Inches(5.1), Inches(4.35))

            _set_notes(slide, slide_content.notes or "")

        output = Path(output_path)
        output.parent.mkdir(parents=True, exist_ok=True)
        prs.save(str(output))
        return str(output)

    # Default: generate with a fixed built-in layout.
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    for index, slide_content in enumerate(slides):
        layout = prs.slide_layouts[1]
        slide = prs.slides.add_slide(layout)

        title_shape = slide.shapes.title
        title_shape.text = slide_content.title
        _apply_title_style(title_shape.text_frame)

        body = slide.placeholders[1]
        figure_idx = bound_figures.get(index)
        has_bound_figure = figure_idx is not None and 0 <= figure_idx < len(figure_items) and Path(figure_items[figure_idx].image_path).exists()
        body.left = Inches(0.9)
        body.top = Inches(1.8)
        body.width = Inches(5.9) if has_bound_figure else Inches(11.4)
        body.height = Inches(4.9)
        text_frame = body.text_frame
        text_frame.clear()

        for bullet_index, bullet in enumerate(slide_content.bullets):
            paragraph = text_frame.paragraphs[0] if bullet_index == 0 else text_frame.add_paragraph()
            paragraph.text = bullet
            paragraph.level = 0
            paragraph.font.size = Pt(21)
            paragraph.font.name = "Arial"
            paragraph.font.color.rgb = BODY_COLOR
            paragraph.space_after = Pt(6)

        # A small accent line helps make the default template look less bare.
        line = slide.shapes.add_shape(1, Inches(0.8), Inches(1.35), Inches(2.4), Inches(0.08))
        fill = line.fill
        fill.solid()
        fill.fore_color.rgb = accent_color
        line.line.color.rgb = accent_color

        if index == 0:
            subtitle = slide.shapes.add_textbox(Inches(0.9), Inches(5.9), Inches(11.0), Inches(0.6))
            subtitle_tf = subtitle.text_frame
            subtitle_tf.text = "Auto-generated academic presentation (editable PPTX)"
            for run in subtitle_tf.paragraphs[0].runs:
                run.font.name = "Arial"
                run.font.size = Pt(16)
                run.font.color.rgb = accent_color

        if has_bound_figure:
            image_path = Path(figure_items[figure_idx].image_path)
            _add_picture_contain(slide, image_path, Inches(7.05), Inches(1.5), Inches(5.4), Inches(4.45))

        _apply_body_style(text_frame)
        _set_notes(slide, slide_content.notes or "")

    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output))
    return str(output)
